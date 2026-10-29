"""Source-supported final refinements for validation prose and download objects."""
import re
from openpyxl import load_workbook

def refine(path, source=None):
    wb=load_workbook(path); changes=[]
    original=load_workbook(source) if source else None
    for ws in wb:
        headers={c.value:c.column for c in ws[1]}
        for row in range(2,ws.max_row+1):
            if 'cn_name' not in headers: continue
            nc=ws.cell(row,headers['cn_name']);dc=ws.cell(row,headers['cn_desc'])
            n,d=nc.value or '',dc.value or ''; old=(n,d)
            d=d.replace('在北美和亚洲开展跨区域的入侵活动','开展入侵活动').replace('开源的荷兰政府案例','开源的政府案例')
            d=d.replace('，我们高度确信其代表伊朗政府进行网络间谍活动','，从事网络间谍活动')
            d=d.replace('它针对独联体国家（不包括俄罗斯和英联邦国家）实施了操作性终止开关','它实施了基于地区判断的操作性终止开关')
            d=d.replace('绝对路径会投放导入基目录','绝对路径会绕过导入基目录').replace('JSON 加载其配置。文件。','JSON 文件加载其配置。')
            if 'Chaskiq' in n and original and '/graphql 接口' in str(original[ws.title].cell(row,headers['cn_desc']).value):
                d=d.replace('Chaskiq /api/graphql 存在', 'Chaskiq /graphql 存在')
            if 'Loom for AWS' in n: n=n.replace('Web应用程序漏洞','AI应用程序漏洞')
            if n.startswith('主机命令行'):
                replacements={'通过 wscript.exe 执行无文件 VBScript 有效载荷':'wscript.exe，无文件 VBScript 有效载荷执行','向 CurrentVersion 运行注册表项 添加恶意值':'CurrentVersion 运行注册表项，恶意值添加','通过 WebDAV 发现 Lateral Target 上的网络共享':'WebDAV，横向目标网络共享发现','通过 computerdefaults.exe 绕过用户访问控制':'computerdefaults.exe，用户账户控制绕过','使用 NtlmThief 工具导出凭据':'NtlmThief，凭据导出','使用命令提示符执行 icacls.exe':'命令提示符，icacls.exe 执行','发现系统和网络信息':'系统和网络信息发现','使用 Schtasks 创建计划任务':'Schtasks，计划任务创建','写入包含勒索信的文件并打开它':'勒索信文件写入并打开','使用 SharpKiller 工具绕过 AMSI':'SharpKiller，AMSI 绕过','通过嵌入式 Python 解释器执行 Shellcode':'嵌入式 Python 解释器，Shellcode 执行','使用 NanoDump 工具 Dump Lsass 进程内存':'NanoDump，LSASS 进程内存转储','通过 Masquerade Finger.exe 进程执行远程命令':'伪装 finger.exe 进程，远程命令执行','使用 WorkItemLoadLibrary 加载 DLL':'WorkItemLoadLibrary，DLL 加载','通过 DumpBrowserSecrets 工具收集浏览器数据':'DumpBrowserSecrets，浏览器数据收集','使用 AES 算法加密目录中的多个文件':'AES，目录文件加密','执行崩溃修复活动系统发现':'CrashFix，系统发现','通过 Minidump 从 LSASS 内存 Dump 中提取凭据':'Minidump，LSASS 内存转储凭据提取','通过编译后的 Python 二进制文件更改用户壁纸':'编译后的 Python 二进制文件，用户壁纸更改'}
                for a,b in replacements.items():n=n.replace(a,b)
            if n.startswith('命令与控制'):
                n=n.replace('”','').replace('，泄漏，','，数据泄漏，')
                if '签入' in n and not re.search(r'，/',n):
                    m=re.search(r'winjak\.com(/\w+\.php)',d)
                    if m:n=n.replace('，签入','，'+m[1]+'，签入')
            if n.startswith('恶意文件传输'):
                kind=None
                for evidence,label in [('恶意 DLL 组件','恶意 DLL 组件文件'),('伪装成配置辅助库','恶意木马组件文件'),('ScreenConnect（','远程访问工具'),('恶意 Java 类','恶意 Java 类文件'),('Windows 可执行文件','恶意 Windows 可执行文件'),('64 位 Linux 可执行文件','恶意 64 位 Linux 可执行文件'),('shell 脚本','恶意 Shell 脚本'),('下载脚本','恶意下载脚本'),('网络代理工具','网络代理工具'),('基于 Windows 的工具','Windows 隧道工具'),('加密货币挖矿程序','加密货币挖矿程序')]:
                    if evidence in d:kind=label;break
                if kind and not re.search(r'文件|脚本|工具|程序|DLL|MSI|存档',n.split(' - ',1)[-1]):n=n.replace('，下载','，'+kind+'，下载')
                n=n.replace('，DLL，','，恶意 DLL 组件文件，').replace('，MSI，','，MSI 安装程序包，')
                if 'Malicious AI Models Download' in n:
                    f=re.search(r'"([^"\n]+\.(?:bin|pth|pt|ckpt))"',d)
                    v=re.search(r'变种 #(\d+)',n)
                    if f and v:n=f'恶意文件传输 - 恶意 AI 模型文件，{f[1]}，下载，变种 #{v[1]}'
                    d=d.replace('此验证动作还原了在网络攻击中下载','此验证动作还原了主机尝试下载').replace(' 样本的行为。',' 样本。')
                    d=d.replace('恶意 AI 模型成为供应链攻击新的途径，网络攻击者在一个看起来非常正常的模型中隐藏有效攻击载荷伪装恶意软件，从而在看似可信的人工智能组件提供一个隐藏的后门，这种新的人工智能供应链威胁，既隐蔽又有弹性。','攻击者将恶意载荷隐藏在看似正常的 AI 模型中，通过人工智能组件植入后门，实施供应链攻击。')
                d=re.sub(r'。ref:(https?://\S+?)/?。$',lambda m:'。\n\n请参考：\n'+m[1]+'/',d)
            if ws.title=='Sequences':
                n=n.replace('活动 攻击活动','攻击活动').replace('绅士勒索软件','Gentlemen 勒索软件')
                d=d.replace('绅士勒索软件','Gentlemen 勒索软件').replace('CrashFix ModeloRAT 活动在','CrashFix ModeloRAT 在')
                if 'Malicious AI Models' in n:
                    n='恶意活动场景 - 恶意 AI 模型文件下载攻击活动'
                    d='此验证场景包括了网络攻击中使用的恶意 AI 模型文件各种变种的下载。攻击者将恶意载荷隐藏在看似正常的 AI 模型中，通过人工智能组件植入后门，实施供应链攻击。'
            if (n,d)!=old:
                changes.append({'sheet':ws.title,'row':row,'old':old,'new':[n,d]});nc.value=n;dc.value=d
    wb.save(path);return changes
